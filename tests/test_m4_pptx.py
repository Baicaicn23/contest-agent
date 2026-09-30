"""M4 收尾验收测试：PPT 大纲 → .pptx 真实文件。

- 渲染器直接测：真 python-pptx 写文件、再真打开断言页数与文字；
- SlidesExporter 用假 LLM（返回固定结构化 JSON）测流水线；
- generate 用例集成：假 runner + 真 exporter，验证自动导出与降级。
"""

from pathlib import Path

import pytest
from pptx import Presentation

from contest_agent.application.usecases.generate_material import GenerateMaterial
from contest_agent.infrastructure.slides.exporter import SlidesExporter
from contest_agent.infrastructure.slides.pptx_renderer import render_pptx

SLIDES = [
    {"title": "项目背景", "bullets": ["官方评审看重创新性", "痛点：报名信息分散"]},
    {"title": "方案设计", "bullets": ["三层架构", "自动盯官网"]},
]


# ---------- 渲染器（真文件） ----------


def test_render_pptx_creates_openable_deck(tmp_path: Path) -> None:
    out = render_pptx(SLIDES, tmp_path / "deck.pptx",
                      title="测试杯·参赛答辩", subtitle="测试竞赛")

    assert out.exists()
    deck = Presentation(out)                      # 用 python-pptx 真打开
    assert len(deck.slides.__iter__.__self__._sldIdLst) == 3  # 封面 + 2 页

    texts = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    joined = "\n".join(texts)
    assert "测试杯·参赛答辩" in joined            # 封面主标题
    assert "测试竞赛" in joined                   # 副标题
    assert "项目背景" in joined and "方案设计" in joined
    assert "自动盯官网" in joined                 # 要点逐条在


def test_render_pptx_tolerates_empty_bullets(tmp_path: Path) -> None:
    out = render_pptx([{"title": "只有标题的一页"}], tmp_path / "one.pptx")
    deck = Presentation(out)
    assert len(list(deck.slides)) == 2            # 封面 + 1 页


# ---------- SlidesExporter（假 LLM） ----------


class FakeSlidesLlm:
    def complete_structured(self, system: str, user: str, schema: dict) -> dict:
        assert "大纲" in system or "PPT" in system
        return {
            "title": "测试杯·参赛答辩",
            "subtitle": "测试竞赛",
            "slides": SLIDES,
        }


def test_slides_exporter_pipeline(tmp_path: Path) -> None:
    exporter = SlidesExporter(FakeSlidesLlm())
    info = exporter.export("# 大纲\n- 要点", tmp_path / "out.pptx")

    assert info["slides"] == 2
    assert info["path"].exists()
    deck = Presentation(info["path"])
    assert len(list(deck.slides)) == 3


def test_slides_exporter_propagates_llm_failure(tmp_path: Path) -> None:
    class BadLlm:
        def complete_structured(self, system, user, schema):
            raise ValueError("LLM 返回的不是合法 JSON")

    exporter = SlidesExporter(BadLlm())
    with pytest.raises(ValueError):
        exporter.export("大纲", tmp_path / "bad.pptx")


# ---------- generate 用例集成（假 runner + 真 exporter） ----------


def _generate_with_exporter(tmp_path: Path, exporter):
    """造一个 generate 用例：假 runner 秒回成功，并按技能约定落盘 ppt-outline.md。"""
    from contest_agent.application.harness.agent_factory import AgentOutcome, MaterialTools

    def fake_runner(profile, system_prompt, user_request, tools_builder, max_iters,
                    **kwargs):
        (Path(kwargs["output_dir"]) if "output_dir" in kwargs else tmp_path / "out")
        return AgentOutcome(final_text="大纲已保存为 ppt-outline.md", error=None), MaterialTools(
            toolkit=None, trace=["save_material(ppt-outline.md, 800 字)"])

    # 假 runner 需要先让 ppt-outline.md 存在——直接在输出目录写一份
    (tmp_path / "out").mkdir(exist_ok=True)
    (tmp_path / "out" / "ppt-outline.md").write_text(
        "# 测试杯答辩\n\n## 1、项目背景\n- 要点 A\n- 要点 B\n", encoding="utf-8")

    class FakeSource:
        def list_notices(self, limit=10):
            return []

        def fetch_detail(self, notice):
            return notice

    class FakeRepo:
        def save_if_absent(self, competition):
            return True

        def list_all(self):
            return [type("C", (), {"name": "测试杯", "type": "exam",
                                   "deadline": None, "notice_url": "u1"})()]

    return GenerateMaterial(
        profile=None,
        source=FakeSource(),
        competition_repository=FakeRepo(),
        output_dir=tmp_path / "out",
        runner=fake_runner,
        exporter=exporter,
    )


def test_generate_auto_exports_pptx(tmp_path: Path) -> None:
    usecase = _generate_with_exporter(tmp_path, SlidesExporter(FakeSlidesLlm()))
    result = usecase.execute(skill_name="ppt-outline")

    assert result.success is True
    assert result.pptx_file == "ppt-outline.pptx"
    assert (tmp_path / "out" / "ppt-outline.pptx").exists()


def test_generate_degrades_when_export_fails(tmp_path: Path) -> None:
    class BoomExporter:
        def export(self, markdown_text, out_path):
            raise RuntimeError("渲染炸了")

    usecase = _generate_with_exporter(tmp_path, BoomExporter())
    result = usecase.execute(skill_name="ppt-outline")

    # 降级原则：大纲在手，pptx 失败不毁任务
    assert result.success is True
    assert result.pptx_file is None
    assert "渲染炸了" in result.pptx_hint
