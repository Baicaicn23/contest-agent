"""python-pptx 渲染器（M4 收尾）：结构化幻灯片数组 → 真正的 .pptx 文件。

输入来自 SlidesExporter（LLM 把大纲 Markdown 整理成的结构化 JSON），
本模块只做"排版"：16:9 页面、封面页（主标题+副标题）、每页标题+要点列表。
刻意保持朴素——能稳定打开、标题要点清晰就是参赛大纲的全部要求；
视觉美化是用户拿回家在 PowerPoint 里干的事。
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

# 统一字号：标题 28 / 正文要点 18（参赛投影的可读下限）
TITLE_PT = 28
BODY_PT = 18
COVER_TITLE_PT = 40


def _add_textbox(slide, left, top, width, height, text, pt, bold=False):
    """往幻灯片放一个文本框（python-pptx 没有"纯文字页"布局，全靠文本框拼）。"""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(pt)
    run.font.bold = bold
    return box


def render_pptx(slides: list[dict], out_path: str | Path,
                title: str = "", subtitle: str = "") -> Path:
    """渲染整套幻灯片并保存。slides 每项 {"title": str, "bullets": [str]}。"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)   # 16:9
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]       # 6 号布局 = 全空白页

    # 封面：主标题居中 + 副标题
    cover = prs.slides.add_slide(blank)
    _add_textbox(cover, 1.0, 2.7, 11.3, 1.3,
                 title or "参赛答辩", COVER_TITLE_PT, bold=True)
    if subtitle:
        _add_textbox(cover, 1.0, 4.1, 11.3, 0.8, subtitle, 20)

    # 内容页：标题在顶，要点逐条往下排
    for slide_data in slides:
        page = prs.slides.add_slide(blank)
        page_title = str(slide_data.get("title") or "未命名")
        _add_textbox(page, 0.7, 0.5, 12.0, 0.9, page_title, TITLE_PT, bold=True)

        bullets = [str(b).strip() for b in slide_data.get("bullets", []) if str(b).strip()]
        if bullets:
            body = page.shapes.add_textbox(Inches(0.9), Inches(1.7), Inches(11.5), Inches(5.2))
            frame = body.text_frame
            frame.word_wrap = True
            for index, bullet in enumerate(bullets):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                run = paragraph.add_run()
                run.text = f"• {bullet}"
                run.font.size = Pt(BODY_PT)
                paragraph.space_after = Pt(10)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path
